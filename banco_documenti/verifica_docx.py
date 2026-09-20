# -*- coding: utf-8 -*-
"""Cosa resta del documento dopo che ci si e' passati sopra.

La domanda non e' «il testo e' cambiato» — quella e' facile. La domanda e'
**cosa si e' perso per strada**: un .docx e' uno zip di parti, e chi lo
ricostruisce a partire dal proprio modello butta via quelle che il modello non
conosce. Il tema, per dirne una, e' dove stanno i caratteri e i colori del
documento: perderlo vuol dire che il contratto di qualcuno si apre diverso da
come lo aveva lasciato, e non lo dice nessuno.
"""
import sys
import zipfile
from pathlib import Path

DOVE = Path(__file__).resolve().parent / "documenti"


def racconta(f: Path) -> None:
    import docx
    print(f"--- {f.name}")
    w = docx.Document(f)
    for p in w.paragraphs:
        if not p.text.strip():
            continue
        pezzi = []
        for r in p.runs:
            marche = []
            if r.bold:
                marche.append("grassetto")
            if r.italic:
                marche.append("corsivo")
            try:
                if r.font.color and r.font.color.rgb:
                    marche.append(f"colore {r.font.color.rgb}")
            except Exception:                               # noqa: BLE001
                pass
            if r.font.size:
                marche.append(f"{r.font.size.pt}pt")
            pezzi.append(f"{r.text!r}" + (f"[{','.join(marche)}]" if marche else ""))
        stile = p.style.name if p.style else "(NESSUNO)"
        print(f"  stile={stile!r}: " + " + ".join(pezzi))
    for t in w.tables:
        stile = t.style.name if t.style else "(NESSUNO)"
        print(f"  tabella stile={stile!r}")


def confronta(prima: Path, dopo: Path) -> None:
    a, b = zipfile.ZipFile(prima), zipfile.ZipFile(dopo)
    na = {n for n in a.namelist() if not n.endswith("/")}
    nb = {n for n in b.namelist() if not n.endswith("/")}
    print(f"\n=== {prima.name} -> {dopo.name}")
    print("  parti perse:   ", sorted(na - nb) or "nessuna")
    print("  parti aggiunte:", sorted(nb - na) or "nessuna")
    print("  parti cambiate:", [n for n in sorted(na & nb) if a.read(n) != b.read(n)])
    print(f"  byte: {prima.stat().st_size} -> {dopo.stat().st_size}")


def main() -> int:
    prima = DOVE / "prova.docx"
    if not prima.is_file():
        print("prima lancia `python3 prepara.py`")
        return 2
    try:
        import docx                                          # noqa: F401
    except ImportError:
        print("serve python-docx per guardare cosa resta")
        return 2
    racconta(prima)
    for nome in ("uscita-chirurgia.docx", "uscita-docx-rs.docx"):
        f = DOVE / nome
        if f.is_file():
            racconta(f)
            confronta(prima, f)
        else:
            print(f"\n({nome} non c'e': lancia il banco che lo produce)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
