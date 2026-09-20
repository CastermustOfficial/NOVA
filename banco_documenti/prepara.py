# -*- coding: utf-8 -*-
"""I tre documenti su cui si prova, fatti apposta per essere difficili.

Non si prova su un file di testo con dentro «ciao»: si prova sulle cose che
in un documento vero si rompono. Un PDF **a due colonne** (perche' e' li' che
un estrattore senza posizioni mescola tutto) con dentro una tabella e una
parola su una seconda pagina. Un `.docx` con uno stile di titolo, grassetto e
colore **in mezzo a un paragrafo**, un corsivo, un carattere piu' grande e una
tabella con il suo stile. Un `.xlsx` con formule, un formato percentuale,
un'intestazione in grassetto colorata e un secondo foglio.

Ognuna di quelle cose e' li' perche' e' quella che si perde.
"""
import sys
from pathlib import Path

DOVE = Path(__file__).resolve().parent / "documenti"


def main() -> int:
    try:
        import pymupdf
        import docx
        from docx.shared import Pt, RGBColor
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as e:
        print(f"manca una libreria per costruire i documenti: {e.name}")
        print("  pip install PyMuPDF python-docx openpyxl")
        return 2

    DOVE.mkdir(parents=True, exist_ok=True)

    d = pymupdf.open()
    p = d.new_page()
    p.insert_text((72, 90), "Relazione trimestrale", fontsize=18, fontname="hebo")
    for i, riga in enumerate([
            "Il margine e' sceso al 12% per via dei costi di",
            "trasporto. La clausola di arbitrato resta quella",
            "del contratto del 2024."]):
        p.insert_text((72, 130 + i * 18), riga, fontsize=11)
    for i, riga in enumerate([
            "Colonna destra: nota a margine che in un",
            "PDF sta a destra e nel testo estratto",
            "finisce in mezzo se non si guarda dove sta."]):
        p.insert_text((320, 130 + i * 18), riga, fontsize=11)
    for i, riga in enumerate([("Voce", "2025", "2026"), ("Ricavi", "1.200", "1.410"),
                              ("Costi", "980", "1.240"), ("Margine", "220", "170")]):
        for j, cella in enumerate(riga):
            p.insert_text((72 + j * 120, 240 + i * 22), cella, fontsize=11,
                          fontname="hebo" if i == 0 else "helv")
    d.new_page().insert_text((72, 90), "Pagina due: la parola CHIAVE sta qui.", fontsize=12)
    d.save(DOVE / "prova.pdf")

    w = docx.Document()
    w.add_heading("Contratto di fornitura", level=1)
    par = w.add_paragraph("Le parti convengono che il ")
    r = par.add_run("termine di consegna")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x30, 0x20)
    par.add_run(" sia di trenta giorni.")
    par2 = w.add_paragraph("Secondo paragrafo, con un ")
    r2 = par2.add_run("corsivo")
    r2.italic = True
    par2.add_run(" dentro, e un carattere piu' grande alla fine: ")
    par2.add_run("QUI").font.size = Pt(16)
    w.add_paragraph("Terzo paragrafo, quello che verra' modificato.")
    t = w.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "Voce"
    t.cell(0, 1).text = "Valore"
    t.cell(1, 0).text = "Penale"
    t.cell(1, 1).text = "2% al giorno"
    w.save(DOVE / "prova.docx")

    wb = Workbook()
    s = wb.active
    s.title = "Conti"
    for c, v in zip("ABCD", ["Voce", "2025", "2026", "Delta"]):
        s[f"{c}1"] = v
        s[f"{c}1"].font = Font(bold=True)
        s[f"{c}1"].fill = PatternFill("solid", start_color="FFE8D8")
    for i, (v, a, b) in enumerate([("Ricavi", 1200, 1410), ("Costi", 980, 1240)], start=2):
        s[f"A{i}"], s[f"B{i}"], s[f"C{i}"] = v, a, b
        s[f"D{i}"] = f"=C{i}-B{i}"
    s["A4"], s["B4"], s["C4"], s["D4"] = "Margine", "=B2-B3", "=C2-C3", "=C4-B4"
    s["B5"] = 0.125
    s["B5"].number_format = "0.0%"
    wb.create_sheet("Note")["A1"] = "Un secondo foglio, che non va perso."
    wb.save(DOVE / "prova.xlsx")

    print(f"fatti in {DOVE}:")
    for f in sorted(DOVE.glob("prova.*")):
        print(f"  {f.name}  {f.stat().st_size} byte")
    return 0


if __name__ == "__main__":
    sys.exit(main())
