# -*- coding: utf-8 -*-
"""Le proposte di NOVA nell'harness, dal demone (D339).

`harness_proponi` (Python) scrive una proposta accanto alle sessioni; la
finestra dell'harness la guarda e la applica con le capacita' `harness.*`
del demone. Qui si prova quel giro sul disco vero, con un progetto finto
che ha le sue prove:

1. la proposta si vede com'e' e come sarebbe;
2. applicata **e provata**, se rompe una prova che passava il file torna
   com'era e la proposta resta;
3. applicata col testo ritoccato da chi guarda, il file cambia, la copia
   `.prima` c'e', la proposta se ne va e la sessione ha i blocchi nuovi;
4. buttata, se ne va senza toccare il file;
5. sono della persona: un modello non le vede e non le puo' chiamare;
6. su un Word e su un PDF il demone scrive **come il Python** — le stesse
   modifiche applicate dalle due parti a due copie dello stesso file danno
   lo stesso documento — e una voce che non vale piu' si ferma (D342).

Esce 2 — «qui non si puo' provare» — se il demone non e' costruito.
"""
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

RADICE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RADICE))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

NOME = "novad.exe" if os.name == "nt" else "novad"
DEMONE = next((p for p in (RADICE / "core" / "target" / "release" / NOME,
                           RADICE / "core" / "target" / "debug" / NOME)
               if p.is_file()), None)
if DEMONE is None:
    print("Il demone non e' costruito per questo sistema. Per averlo:")
    print("  cd core && cargo build --release -p novad")
    sys.exit(2)

passati = 0
falliti: list[str] = []


def controlla(nome, condizione, dettaglio=""):
    global passati
    if condizione:
        passati += 1
        print(f"  [ok ] {nome}")
    else:
        falliti.append(nome)
        print(f"  [NO ] {nome}  {dettaglio}")


from nova.core_client import CoreClient                           # noqa: E402
from nova import harness, harness_modifica                        # noqa: E402

casa = tempfile.mkdtemp(prefix="nova-harness-d-")
nova_dir = Path(casa) / "NOVA"
nova_dir.mkdir(parents=True, exist_ok=True)
# La memoria spenta: questa prova non la usa, e il demone acceso la
# creerebbe e la seminerebbe nel vault del progetto (D365).
(nova_dir / "config.json").write_text('{"kb": {"enabled": false}}', encoding="utf-8")

# Il progetto: un modulo e una prova che lo usa, nella convenzione di NOVA.
progetto = Path(casa) / "progetto"
(progetto / "prove").mkdir(parents=True)
modulo = progetto / "conti.py"
modulo.write_bytes(b"def somma(a, b):\r\n    return a + b\r\n")
(progetto / "prove" / "test_conti.py").write_text(
    "import sys\nsys.path.insert(0, '.')\nfrom conti import somma\n"
    "sys.exit(0 if somma(2, 3) == 5 else 1)\n", encoding="utf-8")

# La sessione e la proposta come le scrive il Python, con le sue funzioni:
# la cartella dell'harness e' la stessa per le due meta'.
os.environ["APPDATA"] = casa
os.environ["HOME"] = casa
# La finestra dell'harness risulta viva (e' questo processo): cosi' il
# Python non prova ad accenderne una.
(nova_dir / "harness").mkdir(parents=True, exist_ok=True)
(nova_dir / "harness" / "finestra.json").write_text(json.dumps({"pid": os.getpid()}), encoding="utf-8")
aperta = harness.apri(str(modulo), radice=str(progetto))
assert aperta.get("ok"), aperta


def proponi(testo):
    d = harness_modifica.proponi(
        [{"azione": "sostituisci", "blocco": "r1", "testo": testo}],
        motivo="prova")
    assert d.get("ok"), d


proponi("    return a - b")

endpoint = (rf"\\.\pipe\nova-harness-{os.getpid()}" if os.name == "nt"
            else str(Path(casa) / "nova.sock"))
ambiente = dict(os.environ)
ambiente.update({"APPDATA": casa, "HOME": casa, "USERPROFILE": casa,
                 "XDG_CONFIG_HOME": casa, "XDG_RUNTIME_DIR": casa,
                 "NOVA_PYTHON": sys.executable})
processo = subprocess.Popen(
    [str(DEMONE), "--endpoint", endpoint, "--log", "warn"],
    env=ambiente, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)

try:
    scadenza = time.time() + 20
    while time.time() < scadenza and processo.poll() is None:
        if CoreClient.disponibile(endpoint):
            break
        time.sleep(0.3)
    else:
        fine = (processo.stderr.read() or b"").decode("utf-8", "replace")[-400:]
        print(f"  il demone non ha risposto su {endpoint}: {fine}")
        processo.kill()
        sys.exit(2)

    with CoreClient(endpoint, timeout=120) as c:
        print("\n1. la proposta si vede com'e' e come sarebbe")
        tutte = c.call("finestra.proposte")["proposte"]
        p = tutte[0] if tutte else {}
        controlla("c'e' una proposta, sul file giusto",
                  len(tutte) == 1 and Path(p.get("file", "")).name == "conti.py", str(tutte)[:200])
        controlla("col testo com'e', a capo di Windows compresi",
                  p.get("originale") == "def somma(a, b):\r\n    return a + b\r\n", repr(p.get("originale")))
        controlla("e come sarebbe, con gli stessi a capo",
                  p.get("proposto") == "def somma(a, b):\r\n    return a - b\r\n", repr(p.get("proposto")))
        controlla("una riga arriva e una se ne va",
                  (p.get("piu"), p.get("meno")) == (1, 1), f"{p.get('piu')} {p.get('meno')}")

        print("\n2. applicata e provata, se rompe una prova non resta")
        # La data del sorgente ad adesso: cosi' la modifica rotta arriva
        # quasi sempre nello stesso secondo in cui la prova di prima l'ha
        # compilato, che e' il caso in cui Python eseguiva la copia vecchia.
        os.utime(modulo, None)
        p["modificato"] = c.call("finestra.proposte")["proposte"][0]["modificato"]
        r = c.call("finestra.applica", {"modifiche": [{"file": p["file"], "atteso": p["modificato"]}],
                                       "verifica": True, "cartella": str(progetto)})
        controlla("non applicata, e lo dice coi test", r.get("ok") is False and r.get("verificato")
                  and "cade quello che prima passava" in r.get("motivo", ""), str(r)[:300])
        controlla("il file e' com'era", modulo.read_bytes() == b"def somma(a, b):\r\n    return a + b\r\n",
                  repr(modulo.read_bytes()))
        controlla("e la proposta e' ancora li'", len(c.call("finestra.proposte")["proposte"]) == 1)

        print("\n3. ritoccata da chi guarda, e provata")
        ritoccato = "def somma(a, b):\r\n    # ritoccata\r\n    return a + b\r\n"
        r = c.call("finestra.applica", {"modifiche": [{"file": p["file"], "testo": ritoccato}],
                                       "verifica": True, "cartella": str(progetto)})
        controlla("applicata, i test passano come prima",
                  r.get("ok") is True and "uguale" in r.get("verdetto", ""), str(r)[:300])
        controlla("il file e' quello ritoccato", modulo.read_bytes() == ritoccato.encode(),
                  repr(modulo.read_bytes()))
        controlla("la copia di prima sta accanto",
                  (progetto / "conti.py.prima").read_bytes() == b"def somma(a, b):\r\n    return a + b\r\n")
        controlla("la proposta se n'e' andata", c.call("finestra.proposte")["proposte"] == [])
        stato = json.loads((nova_dir / "harness" / f"{aperta['sessione']}.json").read_text(encoding="utf-8"))
        controlla("e la sessione ha i blocchi nuovi, gli stessi del Python",
                  stato["blocchi"] == harness._leggi_documento(modulo), str(stato["blocchi"])[:200])

        print("\n4. buttata, se ne va senza toccare il file")
        proponi("    return 0")
        prima = modulo.read_bytes()
        r = c.call("finestra.scarta", {"file": str(modulo)})
        controlla("scartata", r == {"ok": True, "scartata": True}, str(r))
        controlla("il file non e' cambiato", modulo.read_bytes() == prima)
        controlla("e non c'e' piu' niente in attesa", c.call("finestra.proposte")["proposte"] == [])

        print("\n5. i test del progetto, da soli")
        r = c.call("finestra.prova", {"cartella": str(progetto), "file": str(modulo)})
        controlla("riconosce gli script di prova e li esegue",
                  r.get("provabile") and r.get("banco") == "script"
                  and r.get("passate") == ["prove/test_conti.py"], str(r)[:300])
        r = c.call("finestra.prova", {"cartella": str(Path(casa) / "vuota")})
        controlla("dove non c'e' niente da provare lo dice",
                  r.get("provabile") is False and "non ho riconosciuto" in r.get("motivo", ""), str(r))

        print("\n6. sono bottoni della persona")
        elenco = {t["name"] for t in c.request("tools/list")["tools"]}
        controlla("un modello non le vede",
                  not any(n.startswith("finestra_") for n in elenco), str(sorted(elenco))[:200])
        r = c.request("tools/call", {"name": "finestra_applica", "arguments": {"modifiche": []}})
        controlla("e non le puo' chiamare", bool(r.get("isError")), str(r)[:200])

        print("\n7. un Word, dal demone come dal Python")
        import shutil
        import docx as pydocx
        doc = pydocx.Document()
        doc.add_heading("Contratto", level=1)
        doc.add_paragraph("Il canone e' di 500 euro.")
        doc.add_paragraph("Da togliere.")
        doc.add_paragraph("Firma")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "voce", "importo"
        t.cell(1, 0).text, t.cell(1, 1).text = "canone", "500"
        word_a = progetto / "contratto.docx"
        doc.save(str(word_a))
        word_b = Path(casa) / "gemello.docx"
        shutil.copy2(word_a, word_b)
        modifiche_word = [
            {"azione": "sostituisci", "blocco": "p1", "testo": "Il canone e' di 550 euro."},
            {"azione": "dopo", "blocco": "p0", "testo": "Tra le parti."},
            {"azione": "prima", "blocco": "p3", "testo": "Luogo e data."},
            {"azione": "elimina", "blocco": "p2"},
            {"azione": "sostituisci", "blocco": "t0r1", "testo": "canone | 550"},
        ]

        def com_e(f):
            d = pydocx.Document(str(f))
            return ([(x.text, x.style.name) for x in d.paragraphs],
                    [[c.text for c in r.cells] for tb in d.tables for r in tb.rows])

        assert harness.apri(str(word_b)).get("ok")
        assert harness_modifica.proponi(modifiche_word).get("ok")
        py = harness_modifica.applica()
        controlla("il Python applica", py.get("ok"), str(py)[:200])
        aperta_w = harness.apri(str(word_a), radice=str(progetto))
        assert harness_modifica.proponi(modifiche_word, motivo="contratto").get("ok")
        p = next((x for x in c.call("finestra.proposte")["proposte"] if x["file"].endswith("contratto.docx")), {})
        controlla("la proposta si guarda qui, a voci",
                  p.get("qui") is True and p.get("tipo") == "docx" and len(p.get("voci", [])) == 5
                  and not p.get("saltate"), str(p)[:300])
        r = c.call("finestra.applica", {"modifiche": [{"file": p.get("file", ""), "atteso": p.get("modificato")}]})
        controlla("il demone applica", r.get("ok") is True and r["file"][0]["quante"] == 5, str(r)[:300])
        # Gli indici sono quelli del documento com'era: «elimina p2» toglie
        # «Da togliere» anche se prima, nella stessa proposta, si e' aggiunto
        # un paragrafo sopra.
        atteso = ([("Contratto", "Heading 1"), ("Tra le parti.", "Heading 1"),
                   ("Il canone e' di 550 euro.", "Normal"), ("Luogo e data.", "Normal"),
                   ("Firma", "Normal")],
                  [["voce", "importo"], ["canone", "550"]])
        controlla("il documento e' quello chiesto", com_e(word_a) == atteso, str(com_e(word_a)))
        controlla("e lo stesso del Python", com_e(word_a) == com_e(word_b),
                  f"\n     rust {com_e(word_a)}\n     py   {com_e(word_b)}")
        controlla("la copia di prima sta accanto", (progetto / "contratto.docx.prima").is_file())
        stato = json.loads((nova_dir / "harness" / f"{aperta_w['sessione']}.json").read_text(encoding="utf-8"))
        controlla("e la sessione ha i blocchi nuovi, gli stessi del Python",
                  stato["blocchi"] == harness._leggi_documento(word_a),
                  f"\n     rust {stato['blocchi']}\n     py   {harness._leggi_documento(word_a)}")

        print("\n8. una voce che non vale piu' si ferma, le altre si scelgono")
        harness.apri(str(word_a), radice=str(progetto))
        assert harness_modifica.proponi([
            {"azione": "sostituisci", "blocco": "p2", "testo": "Il canone e' di 600 euro."},
            {"azione": "sostituisci", "blocco": "p0", "testo": "Contratto di locazione"},
        ]).get("ok")
        d = pydocx.Document(str(word_a))
        d.paragraphs[2].runs[0].text = "Il canone e' di 580 euro."
        d.save(str(word_a))
        p = next((x for x in c.call("finestra.proposte")["proposte"] if x["file"].endswith("contratto.docx")), {})
        voci = p.get("voci", [])
        controlla("la voce sul paragrafo cambiato e' spenta, col perche'",
                  len(voci) == 2 and "un'altra cosa" in (voci[0].get("guaio") or "") and voci[1].get("guaio") is None,
                  str(voci)[:300])
        try:
            c.call("finestra.applica", {"modifiche": [{"file": p["file"]}]})
            controlla("applicarla tutta non scrive niente", False, "ha scritto")
        except Exception as e:                                  # noqa: BLE001
            controlla("applicarla tutta non scrive niente",
                      "non si applica piu' per intero" in str(e)
                      and com_e(word_a)[0][2][0] == "Il canone e' di 580 euro.", str(e)[:200])
        r = c.call("finestra.applica", {"modifiche": [{"file": p["file"], "scelte": [1]}]})
        controlla("scelta solo quella buona, passa solo quella",
                  r.get("ok") and com_e(word_a)[0][0][0] == "Contratto di locazione"
                  and com_e(word_a)[0][2][0] == "Il canone e' di 580 euro.", str(com_e(word_a))[:200])

        print("\n9. un PDF, annotato dal demone come dal Python")
        import fitz
        pdf = fitz.open()
        pag = pdf.new_page(width=595, height=842)
        pag.insert_text((72, 100), "Articolo 1. Il canone.", fontsize=12)
        pag.insert_text((72, 300), "Articolo 2. La durata.", fontsize=12)
        pag2 = pdf.new_page(width=595, height=842)
        pag2.insert_text((72, 150), "Allegato ruotato.", fontsize=12)
        pag2.set_rotation(90)
        pdf_a = progetto / "contratto.pdf"
        pdf.save(str(pdf_a))
        pdf.close()
        pdf_b = Path(casa) / "gemello.pdf"
        shutil.copy2(pdf_a, pdf_b)
        segni = [
            {"azione": "evidenzia", "blocco": "p0b0"},
            {"azione": "nota", "blocco": "p0b1", "testo": "Durata: perché tre anni?"},
            {"azione": "evidenzia", "blocco": "p1b0"},
        ]

        def annotazioni(f):
            d = fitz.open(str(f))
            # Dove sta il giallo sono i quattro angoli (`vertices`), non il
            # rettangolo dell'annotazione: PyMuPDF lo allarga di qualche
            # punto per sicurezza, e chi legge colora gli angoli.
            fuori = [(n, a.type[1],
                      [round(c) for v in (a.vertices or []) for c in v] if a.type[1] == "Highlight"
                      else [round(x) for x in a.rect],
                      a.info.get("content", ""), a.info.get("title", ""))
                     for n, pg in enumerate(d) for a in pg.annots()]
            d.close()
            return fuori

        assert harness.apri(str(pdf_b)).get("ok")
        assert harness_modifica.proponi(segni).get("ok")
        py = harness_modifica.applica()
        controlla("il Python annota", py.get("ok"), str(py)[:200])
        prima = pdf_a.read_bytes()
        harness.apri(str(pdf_a), radice=str(progetto))
        assert harness_modifica.proponi(segni).get("ok")
        p = next((x for x in c.call("finestra.proposte")["proposte"] if x["file"].endswith("contratto.pdf")), {})
        controlla("la proposta porta dove stanno i blocchi",
                  p.get("tipo") == "pdf" and all(v.get("riquadro") for v in p.get("voci", [])), str(p)[:300])
        r = c.call("finestra.applica", {"modifiche": [{"file": p.get("file", ""), "atteso": p.get("modificato")}]})
        controlla("il demone annota", r.get("ok") is True and r["file"][0]["quante"] == 3, str(r)[:300])
        controlla("le stesse annotazioni del Python, negli stessi punti",
                  annotazioni(pdf_a) == annotazioni(pdf_b),
                  f"\n     rust {annotazioni(pdf_a)}\n     py   {annotazioni(pdf_b)}")
        controlla("scritte in coda: il file di prima e' intatto dentro quello nuovo",
                  pdf_a.read_bytes().startswith(prima))

finally:
    processo.terminate()
    try:
        processo.wait(timeout=10)
    except subprocess.TimeoutExpired:
        processo.kill()

print(f"\n{passati} controlli passati, {len(falliti)} falliti")
if falliti:
    print("::error::test_demone_harness: " + "; ".join(falliti))
    sys.exit(1)
sys.exit(0)
