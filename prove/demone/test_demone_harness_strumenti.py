# -*- coding: utf-8 -*-
"""Gli strumenti `harness_*` del demone contro quelli del Python (D344).

Le due meta' lavorano sugli stessi documenti, ciascuna nella sua cartella
dell'harness, e si confronta quel che dicono e quel che scrivono sul disco:

1. un documento a pezzi: Markdown, testo, codice, un Word — gli stessi
   blocchi, e le stesse risposte di apri, stato, cerca e leggi;
2. un PDF: lo stesso testo, pagina per pagina (i blocchi del demone non
   vengono da PyMuPDF, e i numeri possono essere diversi);
3. una cartella come progetto: lo stesso albero, lo stesso file da cui si
   parte, e le stesse risposte di cerca_progetto;
4. una proposta: la stessa risposta e lo stesso file, con lo stesso nome;
5. applicare: lo stesso documento dopo, e la copia di prima;
6. provare il progetto, e buttare una proposta;
7. un modello li vede, e chiamarli passa dal suo cancello.

Esce 2 — «qui non si puo' provare» — se il demone non e' costruito o manca
python-docx o PyMuPDF.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import warnings
from pathlib import Path

RADICE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RADICE))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
warnings.filterwarnings("ignore")

NOME = "novad.exe" if os.name == "nt" else "novad"
DEMONE = next((p for p in (RADICE / "core" / "target" / "release" / NOME,
                           RADICE / "core" / "target" / "debug" / NOME)
               if p.is_file()), None)
if DEMONE is None:
    print("Il demone non e' costruito per questo sistema. Per averlo:")
    print("  cd core && cargo build --release -p novad")
    sys.exit(2)
try:
    import docx  # noqa: F401
    import fitz  # noqa: F401
except ImportError as e:
    print(f"manca {e.name}: il Python non saprebbe fare la sua meta'")
    sys.exit(2)

passati = 0
falliti: list[str] = []


def controlla(nome, condizione, dettaglio=""):
    global passati
    if condizione:
        passati += 1
        print(f"  [ok ] {nome}")
    else:
        falliti.append(nome)
        print(f"  [NO ] {nome}  {dettaglio}")


from nova.core_client import CoreClient                           # noqa: E402
from nova import harness as H, harness_modifica as M              # noqa: E402

lavoro = Path(tempfile.mkdtemp(prefix="nova-hstr-"))
casa_py = lavoro / "py"
casa_rs = lavoro / "rs"
for c in (casa_py, casa_rs):
    (c / "NOVA" / "harness").mkdir(parents=True)
    # La memoria spenta: questa prova non la usa, e il demone acceso la
    # creerebbe e la seminerebbe nel vault del progetto (D365).
    (c / "NOVA" / "config.json").write_text('{"kb": {"enabled": false}}', encoding="utf-8")
os.environ["APPDATA"] = str(casa_py)
# Nessuna finestra da accendere, da nessuna delle due parti.
H.apri_se_serve = lambda *a, **k: {"viva": True, "accesa_adesso": False, "motivo": ""}
(casa_rs / "NOVA" / "harness" / "finestra.json").write_text(
    json.dumps({"pid": os.getpid()}), encoding="utf-8")
base_py = casa_py / "NOVA" / "harness"
base_rs = casa_rs / "NOVA" / "harness"

# ------------------------------------------------------------- i documenti
docs = lavoro / "documenti"
docs.mkdir()
(docs / "appunti.md").write_text(
    "# Termodinamica\n\nL'entropia di un sistema isolato non diminuisce.\n"
    "E' il secondo principio.\n\n\n## Esempi\n\nIl ghiaccio che si scioglie.\n",
    encoding="utf-8")
(docs / "vecchio.txt").write_bytes(
    "prima riga\r\nseconda\x0criga\r\n\r\nterzo paragrafo è qui\r\n".encode("utf-8"))
(docs / "conti.py").write_text(
    "\ufeffdef somma(a, b):\n    return a + b\n\n\ndef media(xs):\n"
    "    return somma(*xs) / len(xs)\n", encoding="utf-8")
(docs / "pagina.html").write_text("<html><body>\n<h1>Titolo</h1>\n\n<p>entropia</p>\n</body></html>",
                                  encoding="utf-8")
d = docx.Document()
d.add_heading("Contratto di locazione", level=1)
d.add_paragraph("Il canone e' di 500 euro al mese.")
d.add_paragraph("")
d.add_paragraph("La durata e' di quattro anni.", style="List Bullet")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "voce", "importo"
t.cell(1, 0).text, t.cell(1, 1).text = "canone", "500"
d.save(str(docs / "contratto.docx"))
pdf = fitz.open()
pg = pdf.new_page(width=595, height=842)
pg.insert_text((72, 100), "Capitolo 1. L'entropia.", fontsize=14)
pg.insert_textbox(fitz.Rect(72, 130, 520, 300),
                  "L'entropia di un sistema isolato non diminuisce mai. " * 3, fontsize=11)
pg2 = pdf.new_page(width=595, height=842)
pg2.insert_text((72, 150), "Capitolo 2. Il ghiaccio.", fontsize=14)
pdf.save(str(docs / "libro.pdf"))
pdf.close()

endpoint = (rf"\\.\pipe\nova-hstr-{os.getpid()}" if os.name == "nt"
            else str(lavoro / "nova.sock"))
ambiente = dict(os.environ)
ambiente.update({"APPDATA": str(casa_rs), "HOME": str(casa_rs), "USERPROFILE": str(casa_rs),
                 "XDG_CONFIG_HOME": str(casa_rs), "XDG_RUNTIME_DIR": str(lavoro),
                 "NOVA_PYTHON": sys.executable})
processo = subprocess.Popen([str(DEMONE), "--endpoint", endpoint, "--log", "warn"],
                            env=ambiente, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)


def senza(d: dict, *chiavi) -> dict:
    return {k: v for k, v in d.items() if k not in chiavi}


def sessione(base: Path) -> dict:
    s = json.loads((base / "corrente.json").read_text(encoding="utf-8"))["sessione"]
    return json.loads((base / f"{s}.json").read_text(encoding="utf-8"))


try:
    scadenza = time.time() + 20
    while time.time() < scadenza and processo.poll() is None:
        if CoreClient.disponibile(endpoint):
            break
        time.sleep(0.3)
    else:
        print("  il demone non ha risposto:",
              (processo.stderr.read() or b"").decode("utf-8", "replace")[-400:])
        processo.kill()
        sys.exit(2)

    with CoreClient(endpoint, timeout=120) as c:
        print("\n1. un documento a pezzi, dalle due parti")
        DOMANDE = ["entropia sistema isolato", "ghiacco", "canone mensile", "somma",
                   "di il", "titolo"]
        for nome in ["appunti.md", "vecchio.txt", "conti.py", "pagina.html", "contratto.docx"]:
            f = str(docs / nome)
            py = H.apri(f)
            rs = c.call("harness.apri", {"percorso": f})
            controlla(f"{nome}: apri dice le stesse cose",
                      senza(py, "sessione", "finestra") == senza(rs, "sessione", "finestra", "nota"),
                      f"\n     py {senza(py, 'sessione', 'finestra')}\n     rs {senza(rs, 'sessione', 'finestra', 'nota')}")
            sp, sr = sessione(base_py), sessione(base_rs)
            controlla(f"{nome}: gli stessi blocchi nella sessione", sp["blocchi"] == sr["blocchi"],
                      f"\n     py {sp['blocchi'][:3]}\n     rs {sr['blocchi'][:3]}")
            controlla(f"{nome}: e il resto della sessione",
                      senza(sp, "sessione", "aperto", "blocchi") == senza(sr, "sessione", "aperto", "blocchi"),
                      f"\n     py {senza(sp, 'sessione', 'aperto', 'blocchi')}\n     rs {senza(sr, 'sessione', 'aperto', 'blocchi')}")
            controlla(f"{nome}: stato", senza(H.stato(), "sessione") == senza(c.call("harness.stato"), "sessione"))
            for q in DOMANDE:
                a, b = H.cerca(q, quanti=3), c.call("harness.cerca", {"domanda": q, "quanti": 3})
                if a != b:
                    controlla(f"{nome}: cerca «{q}»", False, f"\n     py {a}\n     rs {b}")
                    break
            else:
                controlla(f"{nome}: cerca, {len(DOMANDE)} domande uguali", True)
            controlla(f"{nome}: e le stesse cose evidenziate",
                      sessione(base_py)["evidenziati"] == sessione(base_rs)["evidenziati"])
            primo = sp["blocchi"][min(2, len(sp["blocchi"]) - 1)]["id"]
            for intorno, quanti in [("", 3), (primo, 1), ("z9", 2)]:
                a = H.leggi(intorno=intorno, blocchi=quanti)
                b = c.call("harness.leggi", {"intorno": intorno, "blocchi": quanti})
                controlla(f"{nome}: leggi intorno a «{intorno}»", a == b, f"\n     py {a}\n     rs {b}")

        print("\n2. un PDF: lo stesso testo, pagina per pagina")
        f = str(docs / "libro.pdf")
        py, rs = H.apri(f), c.call("harness.apri", {"percorso": f})
        controlla("si apre dalle due parti, con le stesse pagine",
                  py.get("ok") and rs.get("ok") and py["pagine"] == rs["pagine"], f"{py} {rs}")
        pezzi = lambda s, n: "".join("".join(b["testo"].split()) for b in s["blocchi"] if b["pagina"] == n)
        sp, sr = sessione(base_py), sessione(base_rs)
        controlla("lo stesso testo su ogni pagina", all(pezzi(sp, n) == pezzi(sr, n) for n in (1, 2)),
                  f"\n     py {[b['testo'] for b in sp['blocchi']]}\n     rs {[b['testo'] for b in sr['blocchi']]}")

        def sopra(r, q):
            return not (r[2] < q[0] or q[2] < r[0] or r[3] < q[1] or q[3] < r[1])
        controlla("ogni blocco del demone sta sopra un blocco di PyMuPDF",
                  all(any(x["pagina"] == y["pagina"] and sopra(x["riquadro"], y["riquadro"])
                          for y in sp["blocchi"]) for x in sr["blocchi"]),
                  f"\n     py {[(b['id'], b['riquadro']) for b in sp['blocchi']]}\n     rs {[(b['id'], b['riquadro']) for b in sr['blocchi']]}")
        rc = c.call("harness.cerca", {"domanda": "ghiaccio", "quanti": 1})
        controlla("e la ricerca trova la pagina giusta",
                  rc["trovati"] and rc["trovati"][0]["pagina"] == 2, str(rc))

        print("\n3. una cartella come progetto")
        prog = lavoro / "progetto"
        (prog / "src").mkdir(parents=True)
        (prog / "node_modules" / "x").mkdir(parents=True)
        (prog / "prove").mkdir()
        (prog / "README.md").write_text("Il progetto dei conti.\n\nSi prova con gli script.\n", encoding="utf-8")
        (prog / "src" / "conti.py").write_text("def somma(a, b):\n    return a + b\n", encoding="utf-8")
        (prog / "src" / "a.b.py").write_text("x = 1\n", encoding="utf-8")
        (prog / "src.txt").write_text("il file accanto alla cartella\n", encoding="utf-8")
        (prog / ".gitignore").write_text("target/\n", encoding="utf-8")
        (prog / "Makefile").write_text("prova:\n\tpython prove/test_conti.py\n", encoding="utf-8")
        (prog / "grosso.md").write_text("x " * 300_000, encoding="utf-8")
        (prog / "node_modules" / "x" / "somma.js").write_text("somma", encoding="utf-8")
        (prog / "prove" / "test_conti.py").write_text(
            "import sys\nsys.path.insert(0, 'src')\nfrom conti import somma\n"
            "sys.exit(0 if somma(2, 3) == 5 else 1)\n", encoding="utf-8")
        shutil.copy(docs / "libro.pdf", prog / "libro.pdf")
        py, rs = H.apri(str(prog)), c.call("harness.apri", {"percorso": str(prog)})
        controlla("si apre dallo stesso file, con lo stesso numero di file",
                  senza(py, "sessione", "finestra") == senza(rs, "sessione", "finestra", "nota"),
                  f"\n     py {senza(py, 'sessione', 'finestra')}\n     rs {senza(rs, 'sessione', 'finestra', 'nota')}")
        sp, sr = sessione(base_py), sessione(base_rs)
        controlla("lo stesso albero", sp["albero"] == sr["albero"], f"\n     py {sp['albero']}\n     rs {sr['albero']}")
        for q in ["somma", "script prova", "entropia sistema", "niente di tutto questo xyzzy", ""]:
            a = H.cerca_progetto(q, quanti=5)
            b = senza(c.call("harness.cerca_progetto", {"domanda": q, "quanti": 5}), "nota")
            # Nei PDF i blocchi non sono quelli di PyMuPDF: si confronta il file.
            pdf_a = [(x["file"], x["pagina"]) for x in a.get("risultati", []) if x["file"].endswith(".pdf")]
            pdf_b = [(x["file"], x["pagina"]) for x in b.get("risultati", []) if x["file"].endswith(".pdf")]
            resto = lambda r: [x for x in r.get("risultati", []) if not x["file"].endswith(".pdf")]
            uguali = (senza(a, "risultati") == senza(b, "risultati") and resto(a) == resto(b)
                      and pdf_a == pdf_b)
            controlla(f"cerca_progetto «{q}»", uguali, f"\n     py {a}\n     rs {b}")
        indici = lambda base: sorted(x.name for x in base.glob("indice-*.json"))
        controlla("l'indice del progetto ha lo stesso nome", indici(base_py) == indici(base_rs) != [],
                  f"{indici(base_py)} {indici(base_rs)}")

        print("\n4. una proposta")
        for nome, modifiche in [
            ("appunti.md", [{"blocco": "r2", "testo": "L'entropia non cala."},
                            {"blocco": "r0", "azione": "dopo", "testo": "Appunti del corso."}]),
            ("contratto.docx", [{"blocco": "p1", "testo": "Il canone e' di 550 euro."},
                                {"blocco": "t0r1", "testo": "canone | 550"}]),
            ("libro.pdf", [{"blocco": "p0b0", "azione": "evidenzia"}]),
            ("appunti.md", [{"blocco": "r99", "testo": "x"}, {"blocco": "r2", "azione": "evidenzia"},
                            {"blocco": "r0", "testo": ""}]),
        ]:
            f = str(docs / nome)
            H.apri(f)
            c.call("harness.apri", {"percorso": f})
            a = M.proponi(modifiche, motivo="prova")
            b = senza(c.call("harness.proponi", {"modifiche": modifiche, "motivo": "prova"}), "nota")
            a2 = senza(a, "sessione", "nota")
            b2 = senza(b, "sessione")
            if nome == "libro.pdf":
                for x in a2.get("anteprima", []) + b2.get("anteprima", []):
                    x.pop("prima", None)
            controlla(f"{nome}: proponi risponde uguale", a2 == b2, f"\n     py {a}\n     rs {b}")
            if a.get("ok"):
                fp = M.file_proposta(f)
                fr = base_rs / fp.name
                controlla(f"{nome}: la proposta ha lo stesso nome", fr.is_file(), f"{fp.name}")
                if fr.is_file():
                    pa = json.loads(fp.read_text(encoding="utf-8"))
                    pb = json.loads(fr.read_text(encoding="utf-8"))
                    togli = ("quando", "sessione")
                    if nome == "libro.pdf":
                        for x in pa["modifiche"] + pb["modifiche"]:
                            for k in ("prima", "riquadro"):
                                x.pop(k, None)
                    controlla(f"{nome}: e lo stesso contenuto", senza(pa, *togli) == senza(pb, *togli),
                              f"\n     py {senza(pa, *togli)}\n     rs {senza(pb, *togli)}")
            M.scarta()
            c.call("harness.scarta")

        print("\n5. applicare")
        for nome, modifiche in [
            ("appunti.md", [{"blocco": "r2", "testo": "L'entropia non cala."}]),
            ("contratto.docx", [{"blocco": "p1", "testo": "Il canone e' di 550 euro."},
                                {"blocco": "p0", "azione": "dopo", "testo": "Tra le parti."}]),
        ]:
            copia_py, copia_rs = lavoro / f"py-{nome}", lavoro / f"rs-{nome}"
            shutil.copy2(docs / nome, copia_py)
            shutil.copy2(docs / nome, copia_rs)
            H.apri(str(copia_py))
            M.proponi(modifiche)
            a = M.applica()
            c.call("harness.apri", {"percorso": str(copia_rs)})
            c.call("harness.proponi", {"modifiche": modifiche})
            b = c.call("harness.applica", {})
            controlla(f"{nome}: applicata dalle due parti", a.get("ok") and b.get("ok"), f"{a} {b}")
            controlla(f"{nome}: con lo stesso numero di modifiche", a.get("applicate") == b.get("applicate"),
                      f"{a.get('applicate')} {b.get('applicate')}")
            if nome.endswith(".docx"):
                testi = lambda p: [x.text for x in docx.Document(str(p)).paragraphs]
                uguale = testi(copia_py) == testi(copia_rs)
            else:
                uguale = copia_py.read_bytes() == copia_rs.read_bytes()
            controlla(f"{nome}: e lo stesso documento dopo", uguale)
            controlla(f"{nome}: la copia di prima e' accanto",
                      Path(b.get("copia_di_prima", "")).is_file(), str(b))
        b = c.call("harness.applica", {})
        controlla("senza proposta lo dice", b == {"ok": False, "motivo": "non c'e' nessuna proposta da applicare"}, str(b))

        print("\n6. provare, e buttare")
        c.call("harness.apri", {"percorso": str(prog)})
        r = c.call("harness.prova", {"file": str(prog / "src" / "conti.py")})
        controlla("prova il progetto aperto", r.get("provabile") and r.get("passate") == ["prove/test_conti.py"],
                  str(r)[:300])
        c.call("harness.apri", {"percorso": str(docs / "appunti.md")})
        r = c.call("harness.prova", {})
        controlla("dove non c'e' niente da provare lo dice",
                  r.get("provabile") is False and "Non ho riconosciuto" in r.get("motivo", ""), str(r)[:300])
        c.call("harness.proponi", {"modifiche": [{"blocco": "r0", "testo": "x"}]})
        controlla("buttata", c.call("harness.scarta") == {"ok": True, "scartata": True})
        controlla("e una seconda volta non c'e' niente", c.call("harness.scarta") == {"ok": True, "scartata": False})

        print("\n7. per un modello")
        elenco = {t["name"] for t in c.request("tools/list")["tools"]}
        nomi = {"harness_apri", "harness_cerca", "harness_leggi", "harness_stato", "harness_cerca_progetto",
                "harness_proponi", "harness_applica", "harness_prova", "harness_scarta"}
        controlla("vede i nove strumenti", nomi <= elenco, str(sorted(nomi - elenco)))
        controlla("e non i bottoni della finestra", not any(n.startswith("finestra_") for n in elenco))
        r = c.request("tools/call", {"name": "harness_stato", "arguments": {}})
        controlla("e li chiama", not r.get("isError") and "appunti.md" in json.dumps(r, ensure_ascii=False),
                  str(r)[:200])
finally:
    processo.terminate()
    try:
        processo.wait(timeout=10)
    except subprocess.TimeoutExpired:
        processo.kill()

print(f"\n{passati} controlli passati, {len(falliti)} falliti")
if falliti:
    print("::error::test_demone_harness_strumenti: " + "; ".join(falliti))
    sys.exit(1)
sys.exit(0)
