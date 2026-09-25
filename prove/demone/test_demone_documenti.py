# -*- coding: utf-8 -*-
"""Leggere un documento, dal demone.

`read_document` diventa `documenti.leggi`. Le letture le confronta col
Python `prove/gemelli/test_documenti_rust.py`; qui si prova il ponte: che
Claude Code e il cervello del demone lo vedano con la descrizione del
Python, che legga un Word e un file di testo, e che gli errori arrivino
come errori.

Esce 2 — «qui non si puo' provare» — se il demone non e' costruito.
"""
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

RADICE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RADICE))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

NOME = "novad.exe" if os.name == "nt" else "novad"
DEMONE = next((p for p in (RADICE / "core" / "target" / "release" / NOME,
                           RADICE / "core" / "target" / "debug" / NOME)
               if p.is_file()), None)
if DEMONE is None:
    print("Il demone non e' costruito per questo sistema. Per averlo:")
    print("  cd core && cargo build --release -p novad")
    sys.exit(2)

passati = 0
falliti: list[str] = []


def controlla(nome, condizione, dettaglio=""):
    global passati
    if condizione:
        passati += 1
        print(f"  [ok ] {nome}")
    else:
        falliti.append(nome)
        print(f"  [NO ] {nome}  {dettaglio}")


from nova.core_client import CoreClient                           # noqa: E402


casa = tempfile.mkdtemp(prefix="nova-documenti-d-")
(Path(casa) / "NOVA").mkdir(parents=True, exist_ok=True)
(Path(casa) / "NOVA" / "config.json").write_text("{}", encoding="utf-8")
(Path(casa) / "note.txt").write_bytes("uno\r\ndue è\r\n".encode("utf-8"))
try:
    from docx import Document
    d = Document()
    d.add_paragraph("Contratto di affitto")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Canone", "800 euro"
    d.save(Path(casa) / "contratto.docx")
    WORD = True
except ImportError:
    WORD = False

endpoint = (rf"\\.\pipe\nova-documenti-{os.getpid()}" if os.name == "nt"
            else str(Path(casa) / "nova.sock"))
ambiente = dict(os.environ)
ambiente.update({"APPDATA": casa, "HOME": casa, "USERPROFILE": casa,
                 "XDG_CONFIG_HOME": casa, "XDG_RUNTIME_DIR": casa})
processo = subprocess.Popen(
    [str(DEMONE), "--endpoint", endpoint, "--log", "warn"],
    env=ambiente, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)


def testo_mcp(c, nome, argomenti):
    r = c.request("tools/call", {"name": nome, "arguments": argomenti})
    return r["content"][0]["text"], bool(r.get("isError"))


try:
    scadenza = time.time() + 20
    while time.time() < scadenza and processo.poll() is None:
        if CoreClient.disponibile(endpoint):
            break
        time.sleep(0.3)
    else:
        fine = (processo.stderr.read() or b"").decode("utf-8", "replace")[-400:]
        print(f"  il demone non ha risposto su {endpoint}: {fine}")
        processo.kill()
        sys.exit(2)

    with CoreClient(endpoint, timeout=60) as c:
        print("\n1. la capacita' c'e', con le parole del Python")
        strumenti = {t["name"]: t for t in c.request("tools/list")["tools"]}
        d = strumenti.get("documenti_leggi", {})
        controlla("documenti_leggi c'e', sicura, con la descrizione di read_document",
                  d.get("description", "").startswith("[safe] Legge il CONTENUTO di un PDF")
                  and set(d.get("inputSchema", {}).get("properties", {})) == {"path", "pagine", "foglio"},
                  repr(d)[:200])
        prova = c.call("documenti.leggi", {"path": "/x/y.pdf", "prova": True})
        controlla("l'anteprima e' quella del Python",
                  prova.get("farei") == "Legge il contenuto di /x/y.pdf", repr(prova))

        print("\n2. legge, e dice quando non puo'")
        detto, err = testo_mcp(c, "documenti_leggi", {"path": str(Path(casa) / "note.txt")})
        controlla("un file di testo, con gli a capo di Windows ripuliti",
                  not err and detto == "uno\ndue è\n", repr(detto))
        if WORD:
            detto, err = testo_mcp(c, "documenti_leggi", {"path": str(Path(casa) / "contratto.docx")})
            controlla("un Word, con la sua tabella",
                      not err and detto == "Contratto di affitto\n\n--- tabella 1 ---\nCanone | 800 euro",
                      repr(detto))
        detto, err = testo_mcp(c, "documenti_leggi", {"path": str(Path(casa) / "manca.pdf")})
        controlla("un file che non c'e' e' un errore", err and detto.endswith("manca.pdf non esiste"),
                  repr(detto))

finally:
    processo.terminate()
    try:
        processo.wait(timeout=10)
    except subprocess.TimeoutExpired:
        processo.kill()

print(f"\n{passati} controlli passati, {len(falliti)} falliti")
if falliti:
    print("::error::test_demone_documenti: " + "; ".join(falliti))
    sys.exit(1)
sys.exit(0)
